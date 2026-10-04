import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def apply_background(slide, prs, bg_path):
    left = top = 0
    width = prs.slide_width
    height = prs.slide_height
    pic = slide.shapes.add_picture(bg_path, left, top, width, height)
    # Move the picture to the back
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)
    
def add_custom_text(slide, text, left, top, width, height, font_size=24, bold=False, color=RGBColor(255, 255, 255)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return txBox

def create_cool_presentation(filename="SANGYAN_NiveshRaksha_Cool_Design.pptx"):
    prs = Presentation()
    
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    bg_image = r"C:\Users\anima\.gemini\antigravity-ide\brain\9ac0f32c-17c5-4451-99f9-fab8abcccda7\ppt_background_1791133498312.jpg"
    workflow_img = r"C:\Users\anima\.gemini\antigravity-ide\brain\9ac0f32c-17c5-4451-99f9-fab8abcccda7\ppt_workflow_1791133516510.jpg"
    tech_img = r"C:\Users\anima\.gemini\antigravity-ide\brain\9ac0f32c-17c5-4451-99f9-fab8abcccda7\ppt_techstack_1791133531880.jpg"
    
    # Slide 1: Cover
    blank_slide_layout = prs.slide_layouts[6] # Blank
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    
    add_custom_text(slide, "[PROJECT NAME]", Inches(1), Inches(2), Inches(11), Inches(1.5), font_size=54, bold=True, color=RGBColor(13, 148, 136))
    add_custom_text(slide, "[ONE-LINE VALUE PROPOSITION]", Inches(1), Inches(3.5), Inches(11), Inches(1), font_size=32)
    add_custom_text(slide, "Team: [TEAM NAME] | Track: [TRACK]", Inches(1), Inches(4.5), Inches(11), Inches(1), font_size=24)
    add_custom_text(slide, "[SPACE FOR YOU TO FILL - ADD LOGOS OR DETAILS HERE]", Inches(1), Inches(5.5), Inches(11), Inches(1), font_size=20, color=RGBColor(200, 200, 200))

    # Slide 2: The Problem
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "The Investor-Safety Problem", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "• Target Audience: First-time retail investors and regional language users.\n"
        "• Pattern: Scammers use guaranteed returns, urgency, and SEBI impersonation.\n"
        "• Gap: Existing platforms focus on trading, lacking a safe space to verify claims.\n"
        "• Statistic: [One verified source or statistic here]"
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(4), font_size=28)

    # Slide 3: Target Users & Differentiators
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Target Users & Differentiators", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "Users:\n"
        "• First-time investors (18-35)\n"
        "• Regional-language users (TA/HI/TE/ML/KN)\n"
        "• Senior citizens facing UPI pressure\n\n"
        "Differentiators:\n"
        "• Explainability by construction (no opaque AI decisions)\n"
        "• Honest uncertainty as UX\n"
        "• Privacy as default (redacted, consent-gated)"
    )
    add_custom_text(slide, content, Inches(1), Inches(1.5), Inches(11), Inches(5), font_size=24)

    # Slide 4: Solution Workflow
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "From Suspicious Message to Safer Action", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    # Add workflow image
    slide.shapes.add_picture(workflow_img, Inches(3.1), Inches(1.5), height=Inches(5.5))

    # Slide 5: Tech Stack
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "System Architecture & Tech Stack", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    slide.shapes.add_picture(tech_img, Inches(1), Inches(1.5), height=Inches(5))
    
    content = (
        "• Frontend: Next.js, Tailwind CSS\n"
        "• Backend: FastAPI, Pydantic\n"
        "• Data: PostgreSQL, ChromaDB\n"
        "• AI: Hugging Face, IndicBERT"
    )
    add_custom_text(slide, content, Inches(7), Inches(2.5), Inches(5), Inches(3), font_size=24)

    # Slide 6: Explainable Detection & Privacy
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Explainable Detection & Privacy", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "Detection Engine:\n"
        "• Deterministic rules for guaranteed returns, urgency, fake regulators.\n"
        "• No opaque LLM in the critical safety path.\n\n"
        "Privacy-First Pipeline:\n"
        "• Local PII redaction (PAN, Aadhaar, Phone, Email) before processing.\n"
        "• No raw data hoarding. 72-hour auto-expiry.\n"
        "• Consent-gated evidence locker."
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(5), font_size=28)

    # Slide 7: Multilingual & Pause Features
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Education, Pause & Multilingual", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "Behavioural Pause:\n"
        "• A 30-second checklist interrupting the psychology of urgency in scams.\n\n"
        "Multilingual Education Hub:\n"
        "• 6+ Languages supported (English, Tamil, Hindi, Telugu, etc.)\n"
        "• Offline fallback capability for low-bandwidth users.\n"
        "• Raksha Guide Chatbot: RAG-based multilingual assistance."
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(5), font_size=28)

    # Slide 8: Security & Threat Model
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Security & Threat Model", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "• Network-Free URL Checker: Eliminates SSRF risks by analyzing strings without opening sockets.\n"
        "• Prompt-Injection Resistance: Deterministic rules bypass LLM vulnerabilities.\n"
        "• Rate Limiting: Strict sliding-window limits per IP.\n"
        "• Secure Logging: Secret-shaped and identifier-shaped values scrubbed from logs."
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(5), font_size=28)

    # Slide 9: Evaluation & Metrics
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Evaluation & Metrics", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "• Evaluated against a fixed synthetic dataset measuring false positives (FP) and false negatives (FN).\n"
        "• Current Status: 10/10 evaluation cases pass in backend suite.\n"
        "• No prohibited advice tokens in API responses.\n"
        "• Provenance-first verification with checksummed audit trails.\n"
        "• Metric Placeholders: [Latency], [Precision by Language]"
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(5), font_size=28)

    # Slide 10: Roadmap & Closing
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide, prs, bg_image)
    add_custom_text(slide, "Roadmap & Next Steps", Inches(1), Inches(0.5), Inches(11), Inches(1), font_size=40, bold=True, color=RGBColor(13, 148, 136))
    content = (
        "1. Live SEBI intermediaries integration (replacing fixtures).\n"
        "2. OCR for screenshot analysis expansion.\n"
        "3. Community-reviewed rule packs.\n"
        "4. Feature-phone SMS/IVR integrations.\n\n"
        "NiveshRaksha: Pause. Verify. Protect.\n"
        "[Contact details / Demo link]"
    )
    add_custom_text(slide, content, Inches(1), Inches(2), Inches(11), Inches(5), font_size=28)

    prs.save(filename)
    print(f"Saved {filename}")

if __name__ == '__main__':
    create_cool_presentation()
