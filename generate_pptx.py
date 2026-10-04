import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation(filename="SANGYAN_NiveshRaksha_Generic_Template.pptx"):
    prs = Presentation()
    
    # 1. Slide 1 - Protected Cover
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "[PROJECT NAME]"
    subtitle.text = "[ONE-LINE VALUE PROPOSITION]\nTeam: [TEAM NAME] | Track: [TRACK]"

    # 2. Slide 2 - Problem Context
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "The investor-safety problem"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Who is affected: [Who is affected?]"
    tf.add_paragraph().text = "Pattern: [What scam or misinformation pattern exists?]"
    tf.add_paragraph().text = "Gap: [Why existing responses are insufficient?]"
    tf.add_paragraph().text = "Statistic: [One verified source or statistic]"

    # 3. Slide 3 - Target Users
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Target Users"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "[First-time investor] - Need: X, Risk: Y, Outcome: Z"
    tf.add_paragraph().text = "[Regional-language user] - Need: X, Risk: Y, Outcome: Z"
    tf.add_paragraph().text = "[Senior citizen or family member] - Need: X, Risk: Y, Outcome: Z"
    tf.add_paragraph().text = "[Awareness volunteer] - Need: X, Risk: Y, Outcome: Z"

    # 4. Slide 4 - Solution Overview
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "From suspicious message to safer action"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Message / URL / Screenshot\n↓\nPrivacy redaction\n↓\nLanguage detection\n↓\nRisk rules + AI assistance\n↓\nOfficial-source verification\n↓\nExplainable result\n↓\nSafe next steps and education"

    # 5. Slide 5 - Product Experience
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Product Experience"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- Scam message analysis\n- Suspicious URL checking\n- Advisor/entity verification\n- Evidence draft\n- Multilingual education\n- Raksha Guide chatbot\n- Behavioural pause checklist"

    # 6. Slide 6 - User Journey
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "User Journey"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "User receives suspicious message\n→ Opens NiveshRaksha\n→ Selects language\n→ Pastes message\n→ Sensitive values are redacted\n→ Warning signs are highlighted\n→ Official sources are displayed\n→ User gets safe next steps\n→ User creates an optional redacted report"

    # 7. Slide 7 - AI/ML Pipeline
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "AI/ML Pipeline"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Input validation\n→ Language identification\n→ Script/transliteration normalisation\n→ Privacy redaction\n→ Deterministic red-flag rules\n→ Multilingual classifier\n→ URL/entity checks\n→ Evidence retrieval\n→ Explanation generation\n→ Human-readable multilingual output\n\nCallout: AI assists classification and explanation; it does not provide investment advice."

    # 8. Slide 8 - Multilingual Capability
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Multilingual Capability"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Supported: English, Hindi, Tamil, Telugu, Kannada, Malayalam, Bengali, Marathi, Gujarati, Odia, Punjabi, Assamese, [Additional validated language]\n\nCapabilities:\n- Script detection\n- Code-mixed input\n- Translation or response language selection\n- Per-language evaluation"

    # 9. Slide 9 - Explainable Detection
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Explainable Detection"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Flags: Guaranteed return, urgency, fake regulator, asking credentials, suspicious link, personal account payment, fake KYC threat.\n\nFormat:\n- What we detected\n- Why it matters\n- What we verified\n- What we could not verify\n- What to do next"

    # 10. Slide 10 - Raksha Guide Chatbot
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Raksha Guide Chatbot"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- Retrieval-augmented answers\n- Official and curated knowledge sources\n- Citations and freshness\n- Multilingual interaction\n- Refusal of buy/sell requests\n- Safe fallback when evidence is missing\n\n[PLACEHOLDER: Insert Avatar]"

    # 11. Slide 11 - System Architecture
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "System Architecture"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Client layer: Next.js / TypeScript / Tailwind / shadcn\nAPI layer: FastAPI / validation / rate limits\nSafety layer: redaction / rules / URL protection / input validation\nAI/ML layer: multilingual classifier / embeddings / RAG\nData layer: PostgreSQL / vector store / short-lived sessions\nSource layer: official sources / source metadata / freshness\nOperations layer: Docker / tests / monitoring / security scans"

    # 12. Slide 12 - Technology Stack
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Technology Stack"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Frontend: Next.js, TypeScript, Tailwind CSS, shadcn/ui\nBackend: FastAPI, Pydantic, SQLAlchemy\nAI/ML: PyTorch, Hugging Face, sentence-transformers, ChromaDB\nEngineering: Docker, Pytest, Playwright, Gitleaks, Semgrep, Trivy"

    # 13. Slide 13 - Data and Privacy
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Data and Privacy"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- Redaction before logs and model calls\n- No passwords, OTPs, or UPI PINs\n- Minimal data retention\n- Consent before saving evidence\n- Delete and export controls\n- Synthetic evaluation data\n- No training on private user messages\n\nFlow: User input → redaction → analysis → temporary result → optional redacted draft → automatic expiry/deletion"

    # 14. Slide 14 - Security and Threat Model
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Security and Threat Model"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- SSRF protection\n- Private-IP blocking\n- Redirect limits\n- Safe upload validation\n- Prompt-injection resistance\n- Secret scanning\n- Dependency scanning\n- Rate limiting\n- Source verification\n- Secure error handling"

    # 15. Slide 15 - UI/UX Design Language
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "UI/UX Design Language"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Design System:\n- Palette: slate-50, teal-600, amber-500, red-600\n- Typography: Geist Sans\n- Glass cards, Risk badges\n- Evidence timeline, Source-citation chips\n\nMotion:\n- 150-220ms tab/button transitions\n- Reduced-motion support\n- No flashing alerts"

    # 16. Slide 16 - Repository and Agent Workflow
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Repository and Agent Workflow"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Source repositories\n→ license and security audit\n→ selective reuse decision\n→ architecture planning\n→ UI/UX implementation\n→ backend and AI/ML integration\n→ tests and scans\n→ demo-ready commit"

    # 17. Slide 17 - Evaluation
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Evaluation"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Metrics:\n- [Precision by language]\n- [Recall by language]\n- [F1 score by language]\n- [False-positive rate]\n- [Latency]\n- [Citation coverage]\n- [Accessibility result]\n- [Security scan status]\n\nNote: All metrics must come from reproducible evaluation, not model-card claims alone."

    # 18. Slide 18 - Demo Scenario
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Demo Scenario"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "[Synthetic scam message in selected language]\n\n1. Select language.\n2. Paste message.\n3. Redaction preview.\n4. Highlighted warning signs.\n5. Explanation and sources.\n6. Raksha Guide response.\n7. Pause checklist.\n8. Redacted report draft."

    # 19. Slide 19 - Impact and Scalability
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Impact and Scalability"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Target: [Target user group]\nOutcome: [Expected safety outcome]\nReach: [Languages and regions]\nStrategy: [Low-bandwidth strategy]\nPartners: [Future deployment partners, if verified]\nScale: [Scalability plan]"

    # 20. Slide 20 - Limitations and Responsible AI
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Limitations and Responsible AI"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- A low-risk result does not prove legitimacy.\n- Source availability can change.\n- Models may make mistakes across languages and dialects.\n- This is safety information, not investment advice.\n- Users should independently verify and use official reporting channels."

    # 21. Slide 21 - Roadmap
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Roadmap"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "- Prototype.\n- Multilingual evaluation.\n- Official-source adapters.\n- Accessibility expansion.\n- Privacy and security hardening.\n- Controlled pilot.\n- Monitoring and feedback."

    # 22. Slide 22 - Closing / Contact
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "[One-line closing message]"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Team: [Team members]\nContact: [Contact]\nDemo: [Demo URL]\n\n[QR code placeholder]"

    prs.save(filename)
    print(f"Saved {filename}")

if __name__ == '__main__':
    create_presentation()
